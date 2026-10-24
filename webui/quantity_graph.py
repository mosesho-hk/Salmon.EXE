"""Weekly quantity data and a Matplotlib SVG for the web dashboard."""

import json
import re
from datetime import date, timedelta
from io import StringIO
from threading import Lock

from matplotlib import dates as mdates
from matplotlib.figure import Figure
from openpyxl import load_workbook

RENDER_LOCK = Lock()


def load_quantities(root):
    """Read historical spreadsheet quantities and the saved forecast."""
    workbook = load_workbook(root / "data" / "Salmon_data.xlsx", read_only=True, data_only=True)
    history = []
    # Column B contains ISO week labels (for example, 2026U01).
    # Column C contains the weekly export quantity in tonnes.
    for row in workbook["Vekt"].iter_rows(values_only=True):
        week = re.fullmatch(r"(\d{4})U(\d{2})", str(row[1]).strip())
        # Select observation rows, excluding spreadsheet headings and empty cells.
        if not week or row[2] is None:
            continue
        quantity = float(row[2])
        day = date.fromisocalendar(int(week[1]), int(week[2]), 1)
        history.append({"date": day.isoformat(), "quantity": quantity, "kind": "Historical"})
    workbook.close()

    # Forecast values already have calendar dates in YYYY-MM-DD format.
    forecast = []
    for row in json.loads((root / "forecast.json").read_text(encoding="utf-8")):
        day = date.fromisoformat(row["Fweek"])
        quantity = float(row["Fquantity"])
        forecast.append({"date": day.isoformat(), "quantity": quantity, "kind": "Forecast"})
    # Keep both sets of observations in chronological order.
    for series in (history, forecast):
        series.sort(key=lambda row: row["date"])
    return history, forecast


def build_quantity_graph(root, start=None, end=None):
    """Draw the selected dates and return the graph plus browser hover data."""
    history, forecast = load_quantities(root)
    first = date.fromisoformat(history[0]["date"])
    last = date.fromisoformat(forecast[-1]["date"])
    boundary = date.fromisoformat(forecast[0]["date"])
    end = end or last
    # Initially display the two years ending at the final forecast week.
    if start is None:
        start = end.replace(year=end.year - 2)
    visible = [row for row in history + forecast if start.isoformat() <= row["date"] <= end.isoformat()]

    # Matplotlib rendering is serialized because its shared internals are not thread-safe.
    with RENDER_LOCK:
        figure = Figure(figsize=(12, 5), dpi=72)
        ax = figure.subplots()
        figure.subplots_adjust(left=0.10, right=0.97, bottom=0.18, top=0.88)
        # Use a blue-grey background for history and orange for the forecast period.
        ax.set_facecolor("#f3f6fa")
        ax.axvspan(boundary - timedelta(days=3.5), last + timedelta(days=3.5), color="#fff0da", zorder=0)
        # Draw history as a solid blue line and forecasts as a dashed orange line.
        for series, color, style, label in ((history, "#2463a6", "-", "Historical"), (forecast, "#b85a0d", "--", "Forecast")):
            selected = [row for row in series if start.isoformat() <= row["date"] <= end.isoformat()]
            # Insert gaps rather than drawing through missing weeks.
            xs, ys = [], []
            for row in selected:
                day = date.fromisoformat(row["date"])
                if xs and (day - xs[-1]).days > 7:
                    xs.append(day - timedelta(days=1))
                    ys.append(float("nan"))
                xs.append(day)
                ys.append(row["quantity"])
            ax.plot(xs, ys, color=color, linestyle=style, linewidth=1.7, marker="." if len(selected) < 150 else None, label=label)
        # Set the visible dates, quantity axis, date labels, and legend.
        ax.set_xlim(start - timedelta(days=3), end + timedelta(days=3))
        ax.set_ylim(bottom=0)
        ax.set_ylabel("Weekly export quantity (tonnes)")
        locator = mdates.AutoDateLocator(minticks=4, maxticks=8)
        ax.xaxis.set_major_locator(locator)
        ax.xaxis.set_major_formatter(mdates.ConciseDateFormatter(locator))
        ax.grid(axis="y", alpha=0.2)
        ax.legend(loc="lower left", bbox_to_anchor=(0, 1.01), frameon=False, ncol=2)
        for side in ("top", "right"):
            ax.spines[side].set_visible(False)
        # Save the drawing in memory as SVG so the browser can display it.
        output = StringIO()
        figure.savefig(output, format="svg")
        # Convert each observation to SVG coordinates for the hover marker.
        # SVG measures y from the top; Matplotlib measures it from the bottom.
        # The SVG is 360 points high (5 inches multiplied by 72 points per inch).
        points = []
        for row in visible:
            x, y = ax.transData.transform((mdates.date2num(date.fromisoformat(row["date"])), row["quantity"]))
            points.append({**row, "x": float(x), "y": float(360 - y)})
        svg = output.getvalue()
        figure.clear()
    # Supply the image, hover points, forecast table, and scrollbar date limits.
    return {"svg": svg[svg.index("<svg"):], "points": points, "forecast": forecast,
            "min": first.isoformat(), "max": last.isoformat(), "start": start.isoformat(),
            "end": end.isoformat(), "forecast_start": boundary.isoformat(), "unit": "tonnes"}

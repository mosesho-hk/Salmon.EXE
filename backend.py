"""Local recommendation prototype: python -m uvicorn backend:app --reload."""

import json
from datetime import date
from pathlib import Path

from fastapi import FastAPI, HTTPException, Query
from fastapi.responses import FileResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from pydantic import BaseModel, Field, ValidationError

from optimisation.optimise_algo import optimize_purchasing_schedule

ROOT = Path(__file__).resolve().parent
app = FastAPI(title="Salmon.EXE prototype")


class RecommendationRequest(BaseModel):
    maximum_budget: float = Field(ge=0, allow_inf_nan=False)
    buyer_demand_kg: float = Field(gt=0, multiple_of=0.01, allow_inf_nan=False)
    deadline_weeks: int = Field(ge=1, le=52, strict=True)


class ForecastWeek(BaseModel):
    Fprice: float = Field(ge=0, allow_inf_nan=False)
    Fquantity: float = Field(ge=0, allow_inf_nan=False)
    Fweek: date


@app.get("/")
def home():
    return RedirectResponse("/reccom_eng.html")


@app.post("/recommend")
def recommend(request: RecommendationRequest):
    try:
        raw = json.loads((ROOT / "forecast.json").read_text(encoding="utf-8"))
        forecasts = sorted((ForecastWeek.model_validate(row) for row in raw), key=lambda row: row.Fweek)
        if not forecasts or len({row.Fweek for row in forecasts}) != len(forecasts):
            raise ValueError("Forecast weeks must be non-empty and unique.")
        if any((b.Fweek - a.Fweek).days != 7 for a, b in zip(forecasts, forecasts[1:])):
            raise ValueError("Forecast dates must be consecutive weeks.")
    except (OSError, ValueError, TypeError, ValidationError) as error:
        raise HTTPException(503, "A valid forecast.json is required. Generate the forecast first.") from error
    if request.deadline_weeks > len(forecasts):
        raise HTTPException(422, f"Only {len(forecasts)} forecast weeks are available.")

    # Independent synthetic stock per supplier per week, NOT national export volume.
    # Prices are simulated quotes using the market forecast as their reference.
    suppliers = [("Supplier A", 1.00, 40.0), ("Supplier B", 1.03, 30.0), ("Supplier C", 1.06, 50.0)]
    offers = [
        {"supplier": name, "week": row.Fweek.isoformat(),
         "price_nok_kg": round(row.Fprice * multiplier, 2), "stock_kg": stock}
        for row in forecasts[:request.deadline_weeks]
        for name, multiplier, stock in suppliers
    ]
    try:
        schedule = optimize_purchasing_schedule(
            request.maximum_budget, request.buyer_demand_kg, len(offers),
            [offer["price_nok_kg"] for offer in offers],
            [offer["stock_kg"] for offer in offers],
            period_labels=[str(index) for index in range(len(offers))],
        )
    except ValueError as error:
        raise HTTPException(422, str(error)) from error
    except RuntimeError as error:
        raise HTTPException(503, str(error)) from error
    orders = [
        {"supplier": offers[int(index)]["supplier"], "week": offers[int(index)]["week"],
         "quantity_kg": quantity, "cost_nok": cost,
         "price_nok_kg": offers[int(index)]["price_nok_kg"]}
        for index, cost, quantity in schedule
    ]
    return {
        "recommended_orders": orders,
        "total_cost_nok": round(sum(order["cost_nok"] for order in orders), 2),
        "simulated_supplier_data": True,
        "assumptions": "Simulated weekly supplier quotes and stock; no delivery or storage costs. Deadline is measured from the first forecast week.",
    }


@app.get("/api/recommend")
def recommend_read_only(
    maximum_budget: float = Query(ge=0, allow_inf_nan=False),
    buyer_demand_kg: float = Query(gt=0, multiple_of=0.01, allow_inf_nan=False),
    deadline_weeks: int = Query(ge=1, le=52),
):
    """Read-only calculation for clients whose preview server rejects POST."""
    return recommend(RecommendationRequest(
        maximum_budget=maximum_budget,
        buyer_demand_kg=buyer_demand_kg,
        deadline_weeks=deadline_weeks,
    ))


@app.get("/api/quantity-graph")
def quantity_graph(start: date | None = None, end: date | None = None):
    from webui.quantity_graph import build_quantity_graph

    try:
        return build_quantity_graph(ROOT, start, end)
    except (OSError, KeyError, TypeError, ValueError) as error:
        raise HTTPException(422, f"Unable to load quantity graph: {error}") from error


@app.get("/api/interactive-graph-data")
def interactive_graph_data():
    try:
        workbook = load_workbook(ROOT / "data" / "Salmon_data.xlsx", read_only=True, data_only=True)
        try:
            worksheet = workbook["Vekt"]
            historical_dates = []
            historical_prices = []
            historical_quantities = []
            for row in worksheet.iter_rows(min_row=4, values_only=True):
                week_code, quantity, price = row[1], row[2], row[3]
                if not isinstance(week_code, str) or "U" not in week_code:
                    continue
                year, week = (int(part) for part in week_code.split("U", 1))
                week_date = date.fromisocalendar(year, week, 1)
                historical_dates.append(week_date.isoformat())
                historical_quantities.append(float(quantity))
                historical_prices.append(round(float(price) * 0.08, 2))
        finally:
            workbook.close()

        forecast_rows = json.loads((ROOT / "forecast.json").read_text(encoding="utf-8"))
        forecasts = [ForecastWeek.model_validate(row) for row in forecast_rows]
        if not historical_dates or not forecasts:
            raise ValueError("Historical and forecast data must both contain observations.")
        return {
            "historical": {
                "dates": historical_dates,
                "price_gbp": historical_prices,
                "quantity_tonnes": historical_quantities,
            },
            "forecast": {
                "dates": [row.Fweek.isoformat() for row in forecasts],
                "price_gbp": [round(row.Fprice * 0.08, 2) for row in forecasts],
                "quantity_tonnes": [row.Fquantity for row in forecasts],
            },
        }
    except (OSError, KeyError, TypeError, ValueError, ValidationError) as error:
        raise HTTPException(503, f"Unable to load interactive graph data: {error}") from error


@app.get("/data/{filename}")
def graph_image(filename: str):
    # Expose only the two chart images, not the source data directory.
    if filename not in {"Figure_1.png", "Figure_2.png"}:
        raise HTTPException(404, "Graph not found.")
    path = ROOT / "data" / filename
    if not path.is_file():
        raise HTTPException(404, "Graph image has not been generated yet.")
    return FileResponse(path, media_type="image/png")


# Serve the existing frontend and API on the same origin.
app.mount("/", StaticFiles(directory=ROOT / "src", html=True), name="frontend")
